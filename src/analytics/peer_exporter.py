"""Excel exporter for peer comparison report."""

import logging
import sqlite3
import os
from typing import Dict, Any, List
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
import statistics

logger = logging.getLogger(__name__)

# 10 metrics from Day 18
METRICS = [
    "ROE",
    "ROCE",
    "Net Profit Margin",
    "D/E",
    "FCF",
    "PAT CAGR 5yr",
    "Revenue CAGR 5yr",
    "EPS CAGR 5yr",
    "Interest Coverage",
    "Asset Turnover",
]


def generate_peer_comparison_excel(
    connection: sqlite3.Connection, output_path: str = "output/peer_comparison.xlsx"
) -> None:
    """Generate an Excel workbook with peer group comparisons."""
    os.makedirs(os.path.dirname(output_path), exist_ok=True)

    # 1. Get available peer groups
    # We will use peer_percentiles table as it contains the grouped data
    try:
        latest_year = connection.execute(
            "SELECT MAX(year) FROM peer_percentiles"
        ).fetchone()[0]
    except sqlite3.OperationalError:
        logger.error("Table peer_percentiles not found.")
        return

    if latest_year is None:
        logger.warning("No data in peer_percentiles.")
        return

    peer_groups = [
        r[0]
        for r in connection.execute(
            "SELECT DISTINCT peer_group_name FROM peer_percentiles WHERE year=?",
            (latest_year,),
        ).fetchall()
    ]

    if not peer_groups:
        logger.warning("No peer groups found in peer_percentiles.")
        return

    # Fetch company names
    company_names = {
        r[0]: r[1]
        for r in connection.execute(
            "SELECT company_id, company_name FROM companies"
        ).fetchall()
    }

    wb = Workbook()
    wb.remove(wb.active)  # Remove default sheet

    green_fill = PatternFill(
        start_color="c6efce", end_color="c6efce", fill_type="solid"
    )
    yellow_fill = PatternFill(
        start_color="ffeb9c", end_color="ffeb9c", fill_type="solid"
    )
    red_fill = PatternFill(start_color="ffc7ce", end_color="ffc7ce", fill_type="solid")
    gold_fill = PatternFill(
        start_color="ffd700", end_color="ffd700", fill_type="solid"
    )  # Gold/Amber for benchmark

    bold_font = Font(bold=True)
    center_align = Alignment(horizontal="center")

    for pg_name in peer_groups:
        ws = wb.create_sheet(title=pg_name[:31])  # Excel sheet name limit is 31 chars

        # Headers
        headers = ["company_id", "company_name"]
        for m in METRICS:
            headers.append(f"{m} (Raw)")
            headers.append(f"{m} (Rank)")

        ws.append(headers)

        # Style headers
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = bold_font
            cell.alignment = center_align

        # Fetch data for this peer group
        rows = connection.execute(
            """
            SELECT company_id, metric, value, percentile_rank 
            FROM peer_percentiles 
            WHERE peer_group_name=? AND year=?
        """,
            (pg_name, latest_year),
        ).fetchall()

        # Group by company
        company_data: Dict[str, Dict[str, Any]] = {}
        for cid, m, val, rank in rows:
            if cid not in company_data:
                company_data[cid] = {}
            company_data[cid][m] = {"val": val, "rank": rank}

        # Write rows
        start_row = 2

        # We need to compute composite scores to find benchmark (highest composite score)
        # assuming Benchmark is the top company if not explicitly known
        comp_scores = {}
        for cid, m_data in company_data.items():
            ranks = [
                m_data[m]["rank"]
                for m in METRICS
                if m in m_data and m_data[m]["rank"] is not None
            ]
            comp_scores[cid] = sum(ranks) / len(ranks) if ranks else 0.0

        benchmark_cid = (
            max(comp_scores.items(), key=lambda x: x[1])[0] if comp_scores else None
        )

        # Collect values for median calculation
        metric_values: Dict[str, List[float]] = {m: [] for m in METRICS}

        for row_idx, (cid, m_data) in enumerate(company_data.items(), start=start_row):
            row_vals = [cid, company_names.get(cid, "Unknown")]

            for m in METRICS:
                val = m_data.get(m, {}).get("val")
                rank = m_data.get(m, {}).get("rank")

                row_vals.append(val)
                row_vals.append(rank)

                if val is not None:
                    metric_values[m].append(val)

            ws.append(row_vals)

            # Apply styling
            for col_idx, val in enumerate(row_vals, start=1):
                cell = ws.cell(row=row_idx, column=col_idx)

                # Check benchmark
                if cid == benchmark_cid:
                    cell.fill = gold_fill
                else:
                    # Apply percentile coloring if it's a rank column
                    header_name = headers[col_idx - 1]
                    if "(Rank)" in header_name and isinstance(val, (int, float)):
                        if val >= 0.75:
                            cell.fill = green_fill
                        elif val <= 0.25:
                            cell.fill = red_fill
                        else:
                            cell.fill = yellow_fill

        # Summary Row (Median)
        summary_row = ["", "Peer Group Median"]
        for m in METRICS:
            vals = metric_values[m]
            median_val = statistics.median(vals) if vals else None
            summary_row.append(median_val)
            summary_row.append(None)  # No rank for median

        ws.append(summary_row)

        # Style summary row
        last_row = ws.max_row
        for col_idx in range(1, len(summary_row) + 1):
            cell = ws.cell(row=last_row, column=col_idx)
            cell.font = bold_font

        # Auto-adjust column widths
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = max_length + 2
            ws.column_dimensions[column].width = adjusted_width

    wb.save(output_path)
    logger.info("Successfully exported peer comparison report to %s", output_path)


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)
    db_path = "nifty100.db"
    conn = sqlite3.connect(db_path)
    try:
        generate_peer_comparison_excel(conn)
    finally:
        conn.close()
