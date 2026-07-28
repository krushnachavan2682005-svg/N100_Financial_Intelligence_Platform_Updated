import os
import sqlite3
import pytest
from openpyxl import load_workbook
from src.analytics.peer_exporter import generate_peer_comparison_excel

@pytest.fixture
def setup_db(tmp_path):
    conn = sqlite3.connect(":memory:")
    
    conn.executescript("""
        CREATE TABLE companies (company_id TEXT PRIMARY KEY, company_name TEXT);
        CREATE TABLE peer_percentiles (
            company_id TEXT,
            peer_group_name TEXT,
            metric TEXT,
            value REAL,
            percentile_rank REAL,
            year INTEGER
        );
    """)
    
    # 11 peer groups
    peer_groups = [f"Group_{i}" for i in range(1, 12)]
    
    for i, pg in enumerate(peer_groups):
        cid = f"C{i}"
        conn.execute("INSERT INTO companies (company_id, company_name) VALUES (?, ?)", (cid, f"Company {i}"))
        
        # Insert 10 metrics for each
        metrics = ["ROE", "ROCE", "Net Profit Margin", "D/E", "FCF", "PAT CAGR 5yr", "Revenue CAGR 5yr", "EPS CAGR 5yr", "Interest Coverage", "Asset Turnover"]
        for m in metrics:
            conn.execute(
                "INSERT INTO peer_percentiles (company_id, peer_group_name, metric, value, percentile_rank, year) VALUES (?, ?, ?, ?, ?, ?)",
                (cid, pg, m, 10.0, 0.8, 2023)
            )
            
    yield conn
    conn.close()

def test_generate_peer_comparison_excel(setup_db, tmp_path):
    conn = setup_db
    out_file = tmp_path / "peer_comparison.xlsx"
    
    generate_peer_comparison_excel(conn, str(out_file))
    
    assert os.path.exists(out_file)
    
    # Check structure
    wb = load_workbook(out_file)
    assert len(wb.sheetnames) == 11
    
    # Check sheet contents
    ws = wb[wb.sheetnames[0]]
    # 2 columns for company, 10 metrics * 2 columns = 20 columns
    # Total = 22 columns
    assert ws.max_column == 22
    
    # Check header row
    header_vals = [cell.value for cell in ws[1]]
    assert header_vals[0] == "company_id"
    assert header_vals[1] == "company_name"
    assert "ROE (Raw)" in header_vals
    assert "ROE (Rank)" in header_vals
    
    # Row 2 should be the company data (Gold fill since it's the benchmark/only company)
    assert ws.cell(row=2, column=1).value.startswith("C")
    
    # Row 3 should be Summary row
    assert ws.cell(row=3, column=2).value == "Peer Group Median"
    assert ws.cell(row=3, column=3).value == 10.0 # Median of raw value
